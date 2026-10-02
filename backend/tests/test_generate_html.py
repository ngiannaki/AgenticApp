"""generate_file's html format: a styled, self-contained report page."""
from __future__ import annotations

import base64
from types import SimpleNamespace

import pytest

from backend.mcp_servers.code.errors import ValidationError
from backend.mcp_servers.code.generation.generator import generate_file


def test_html_report_with_kpis_chart_and_table() -> None:
    content = {
        "title": "NVIDIA — Balance Sheet <Dashboard>",
        "subtitle": "USD millions",
        "metrics": [{"label": "Total assets", "value": "$111.6B", "note": "+69.8% YoY"}],
        "sections": [{
            "heading": "Assets",
            "paragraphs": ["Strong growth."],
            "chart": {"title": "Assets", "labels": ["FY24", "FY25"], "series": [{"name": "Assets", "values": [65728, 111601]}]},
            "table": {"headers": ["Item", "FY24", "FY25"], "rows": [["Total assets", "65,728", "111,601"]]},
        }],
    }
    out = generate_file("html", "nvidia_dashboard", content, settings=SimpleNamespace(max_artifact_bytes=5_000_000))
    assert out["filename"] == "nvidia_dashboard.html" and out["content_type"] == "text/html"
    page = base64.b64decode(out["data_base64"]).decode()
    assert "&lt;Dashboard&gt;" in page and "<script" not in page  # escaped, no scripts
    assert 'class="kpi"' in page and "+69.8% YoY" in page and 'class="note up"' in page
    assert "<svg" in page and "<rect" in page
    assert '<td class="num">111,601</td>' in page


def test_unknown_format_points_to_generate_ui() -> None:
    with pytest.raises(ValidationError, match="generate_ui"):
        generate_file("pdf", "x", {"title": "t"}, settings=SimpleNamespace(max_artifact_bytes=1000))


def test_pptx_deck_has_title_kpis_chart_and_table_slides() -> None:
    import io

    from pptx import Presentation

    content = {
        "title": "NVIDIA Balance Sheet",
        "subtitle": "USD millions",
        "metrics": [{"label": "Total assets", "value": "$111.6B", "note": "+69.8% YoY"}],
        "sections": [{
            "heading": "Assets",
            "bullets": ["Cash grew strongly", "Inventory up"],
            "chart": {"title": "Assets", "labels": ["FY24", "FY25"], "series": [{"name": "Assets", "values": [65728, 111601]}]},
            "table": {"headers": ["Item", "FY24", "FY25"], "rows": [["Total assets", "65,728", "111,601"]] * 12},
            "notes": "speaker notes",
        }],
    }
    out = generate_file("pptx", "nvidia", content, settings=SimpleNamespace(max_artifact_bytes=5_000_000))
    assert out["filename"] == "nvidia.pptx" and out["content_type"].endswith("presentationml.presentation")
    prs = Presentation(io.BytesIO(base64.b64decode(out["data_base64"])))
    # title, KPIs, assets body (bullets + chart), 12 rows -> 2 table slides
    assert len(prs.slides) == 5
    assert any(sh.has_chart for sh in prs.slides[2].shapes)
    assert sum(sh.has_table for s in prs.slides for sh in s.shapes) == 2
    assert prs.slides[2].notes_slide.notes_text_frame.text == "speaker notes"
