"""generate_file's pptx format: a themed 16:9 slide deck built with python-pptx.

Same document shape as md/docx/html (title + sections, plus the html-only
`subtitle`, `metrics` and per-section `chart`), mapped to slides:

- title (+ subtitle)        -> a dark title slide
- metrics                   -> KPI-card slides (up to 8 cards per slide)
- each section              -> a slide: heading bar, text on the left, a native
                               (editable) chart on the right when present
- a section's table         -> its own slide with a styled native table, split
                               every _ROWS_PER_SLIDE rows
- optional section `notes`  -> speaker notes
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from backend.mcp_servers.code.generation.documents import normalize_document

_NAVY = RGBColor(0x0F, 0x17, 0x2A)
_ACCENT = RGBColor(0x76, 0xB9, 0x00)  # a green accent that reads on navy and on white
_FG = RGBColor(0x1F, 0x29, 0x37)
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_CARD = RGBColor(0xF3, 0xF4, 0xF6)
_LINE = RGBColor(0xE5, 0xE7, 0xEB)
_GOOD = RGBColor(0x16, 0xA3, 0x4A)
_BAD = RGBColor(0xDC, 0x26, 0x26)
_SERIES = [RGBColor(0x25, 0x63, 0xEB), RGBColor(0x76, 0xB9, 0x00), RGBColor(0xF5, 0x9E, 0x0B),
           RGBColor(0x93, 0x33, 0xEA), RGBColor(0x08, 0x91, 0xB2), RGBColor(0xDC, 0x26, 0x26)]
_FONT = "Calibri"

_W, _H = Inches(13.333), Inches(7.5)
_MARGIN = Inches(0.6)
_ROWS_PER_SLIDE = 9
_KPIS_PER_SLIDE = 8


def _text(slide: Any, x: Any, y: Any, w: Any, h: Any, text: str, *, size: int, bold: bool = False,
          color: RGBColor = _FG, align: Any = PP_ALIGN.LEFT, anchor: Any = MSO_ANCHOR.TOP) -> Any:
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, _FONT
    run.font.color.rgb = color
    return box


def _rect(slide: Any, x: Any, y: Any, w: Any, h: Any, fill: RGBColor, shape: Any = MSO_SHAPE.RECTANGLE) -> Any:
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _blank(prs: Any) -> Any:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide: Any, title: str, page: int) -> None:
    _text(slide, _MARGIN, Inches(0.4), _W - 2 * _MARGIN, Inches(0.9), title, size=30, bold=True, color=_NAVY,
          anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, _MARGIN, Inches(1.3), Inches(1.2), Emu(45720), _ACCENT)
    _text(slide, _W - Inches(1.4), _H - Inches(0.5), Inches(0.8), Inches(0.3), str(page), size=11, color=_MUTED,
          align=PP_ALIGN.RIGHT)


def _title_slide(prs: Any, title: str, subtitle: str | None) -> None:
    s = _blank(prs)
    _rect(s, 0, 0, _W, _H, _NAVY)
    _rect(s, _MARGIN, Inches(2.55), Inches(1.6), Emu(64008), _ACCENT)
    _text(s, _MARGIN, Inches(2.8), _W - 2 * _MARGIN, Inches(1.8), title, size=44, bold=True,
          color=RGBColor(0xFF, 0xFF, 0xFF), anchor=MSO_ANCHOR.TOP)
    if subtitle:
        _text(s, _MARGIN, Inches(4.7), _W - 2 * _MARGIN, Inches(1.0), subtitle, size=20,
              color=RGBColor(0xC7, 0xCD, 0xD6))


def _kpi_slides(prs: Any, metrics: list[dict[str, Any]], page: int) -> int:
    for start in range(0, len(metrics), _KPIS_PER_SLIDE):
        chunk = metrics[start:start + _KPIS_PER_SLIDE]
        s = _blank(prs)
        _header(s, "Key metrics" if start == 0 else "Key metrics (cont.)", page)
        page += 1
        cols = 4 if len(chunk) > 3 else len(chunk)
        gap = Inches(0.3)
        card_w = int((_W - 2 * _MARGIN - gap * (cols - 1)) / cols)
        card_h = Inches(2.3)
        for i, m in enumerate(chunk):
            x = _MARGIN + (card_w + gap) * (i % cols)
            y = Inches(1.9) + (card_h + gap) * (i // cols)
            _rect(s, x, y, card_w, card_h, _CARD, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
            _text(s, x + Inches(0.2), y + Inches(0.2), card_w - Inches(0.4), Inches(0.4), str(m.get("label") or "").upper(),
                  size=12, bold=True, color=_MUTED)
            _text(s, x + Inches(0.2), y + Inches(0.7), card_w - Inches(0.4), Inches(0.8), str(m.get("value") or ""),
                  size=32, bold=True, color=_NAVY, anchor=MSO_ANCHOR.MIDDLE)
            note = str(m.get("note") or m.get("change") or "")
            if note:
                trend = {"good": _GOOD, "bad": _BAD}.get(str(m.get("trend") or "").lower())
                if trend is None:
                    n = note.strip()
                    trend = _GOOD if n.startswith("+") else _BAD if n.startswith(("-", "−")) else _MUTED
                _text(s, x + Inches(0.2), y + Inches(1.6), card_w - Inches(0.4), Inches(0.5), note, size=14, color=trend)
    return page


def _chart(slide: Any, chart: dict[str, Any], x: Any, y: Any, w: Any, h: Any) -> bool:
    labels = [str(v) for v in (chart.get("labels") or [])][:24]
    series = [s for s in (chart.get("series") or []) if isinstance(s, dict)][:6]
    if not labels or not series:
        return False
    data = CategoryChartData()
    data.categories = labels
    for s in series:
        vals = [float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0 for v in (s.get("values") or [])]
        data.add_series(str(s.get("name") or ""), (vals + [0.0] * len(labels))[: len(labels)])
    kind = XL_CHART_TYPE.LINE_MARKERS if str(chart.get("type") or "").lower() == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    gframe = slide.shapes.add_chart(kind, x, y, w, h, data)
    ch = gframe.chart
    ch.font.size, ch.font.name = Pt(12), _FONT
    if chart.get("title"):
        ch.has_title = True
        ch.chart_title.text_frame.text = str(chart["title"])
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    else:
        ch.has_title = False
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    ch.value_axis.major_gridlines.format.line.color.rgb = _LINE
    ch.value_axis.format.line.fill.background()
    ch.value_axis.tick_labels.number_format = '#,##0'
    ch.value_axis.tick_labels.number_format_is_linked = False
    for i, plot_series in enumerate(ch.plots[0].series):
        color = _SERIES[i % len(_SERIES)]
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            plot_series.format.line.color.rgb = color
            plot_series.format.line.width = Pt(2.5)
            plot_series.smooth = False
        else:
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = color
    if kind == XL_CHART_TYPE.COLUMN_CLUSTERED and len(labels) <= 8:
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(10)
        plot.data_labels.number_format = '#,##0'
        plot.data_labels.number_format_is_linked = False
    return True


def _body_slide(prs: Any, heading: str, sec: dict[str, Any], chart: Any, page: int) -> Any:
    s = _blank(prs)
    _header(s, heading, page)
    has_chart = isinstance(chart, dict)
    top, height = Inches(1.7), _H - Inches(2.5)
    text_w = Inches(5.6) if has_chart else _W - 2 * _MARGIN

    items: list[tuple[str, str]] = [("p", p) for p in sec["paragraphs"]]
    items += [("b", b) for b in sec["bullets"]]
    items += [("n", f"{i}. {n}") for i, n in enumerate(sec["numbered"], 1)]
    if items:
        total = sum(len(t) for _, t in items)
        size = 20 if total < 350 else 18 if total < 600 else 16 if total < 900 else 14
        box = s.shapes.add_textbox(_MARGIN, top, text_w, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, (kind, t) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            run = p.add_run()
            run.text = ("• " + t) if kind == "b" else t
            run.font.size, run.font.name = Pt(size), _FONT
            run.font.color.rgb = _FG
    drawn = False
    if has_chart:
        cx = _MARGIN + text_w + Inches(0.3) if items else _MARGIN
        cw = _W - _MARGIN - cx
        drawn = _chart(s, chart, cx, top, cw, height)
    return s, drawn


def _table_slides(prs: Any, heading: str, table: dict[str, Any], page: int) -> int:
    headers, rows = table["headers"], table["rows"]
    n_cols = len(headers) if headers else len(rows[0])
    for start in range(0, len(rows), _ROWS_PER_SLIDE):
        chunk = rows[start:start + _ROWS_PER_SLIDE]
        s = _blank(prs)
        _header(s, heading if start == 0 else f"{heading} (cont.)", page)
        page += 1
        n_rows = len(chunk) + (1 if headers else 0)
        row_h = Inches(0.5)
        shape = s.shapes.add_table(n_rows, n_cols, _MARGIN, Inches(1.7), _W - 2 * _MARGIN, row_h * n_rows)
        tbl = shape.table
        # first column wider (labels), the rest share the remainder
        total_w = _W - 2 * _MARGIN
        first = int(total_w * (0.38 if n_cols > 2 else 0.5))
        rest = int((total_w - first) / max(n_cols - 1, 1)) if n_cols > 1 else 0
        tbl.columns[0].width = first if n_cols > 1 else total_w
        for c in range(1, n_cols):
            tbl.columns[c].width = rest
        size = 14 if n_cols <= 5 else 12 if n_cols <= 7 else 10
        numeric = [all(_is_num(r[c]) or r[c] in ("", "—", "-") for r in rows) for c in range(n_cols)]
        grid = ([headers] if headers else []) + chunk
        for r, row in enumerate(grid):
            is_head = bool(headers) and r == 0
            for c, text in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = text
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = _NAVY if is_head else (_CARD if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
                para = cell.text_frame.paragraphs[0]
                para.alignment = PP_ALIGN.RIGHT if numeric[c] and c > 0 else PP_ALIGN.LEFT
                for run in para.runs:
                    run.font.size, run.font.name = Pt(size), _FONT
                    run.font.bold = is_head
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if is_head else _FG
    return page


def _is_num(cell: str) -> bool:
    s = cell.replace(",", "").replace("$", "").replace("%", "").replace("€", "").strip()
    if s.startswith("(") and s.endswith(")"):
        s = s[1:-1]
    try:
        float(s)
        return True
    except ValueError:
        return False


def render_pptx(content: dict[str, Any]) -> bytes:
    doc = normalize_document(content)
    raw_sections = content.get("sections") or []
    prs = Presentation()
    prs.slide_width, prs.slide_height = _W, _H

    subtitle = content.get("subtitle") if isinstance(content.get("subtitle"), str) else None
    if doc["title"]:
        _title_slide(prs, doc["title"], subtitle)
    page = 2
    metrics = [m for m in (content.get("metrics") or []) if isinstance(m, dict)][:16]
    if metrics:
        page = _kpi_slides(prs, metrics, page)

    for i, sec in enumerate(doc["sections"]):
        raw = raw_sections[i] if i < len(raw_sections) and isinstance(raw_sections[i], dict) else {}
        heading = sec["heading"] or doc["title"] or "Overview"
        first_slide = None
        if sec["paragraphs"] or sec["bullets"] or sec["numbered"] or isinstance(raw.get("chart"), dict):
            first_slide, _ = _body_slide(prs, heading, sec, raw.get("chart"), page)
            page += 1
        if sec["table"]:
            n_before = len(prs.slides)
            page = _table_slides(prs, heading if first_slide is None else f"{heading} — data", sec["table"], page)
            first_slide = first_slide or prs.slides[n_before]
        if first_slide is None:  # heading-only section: a divider slide
            first_slide = _blank(prs)
            _header(first_slide, heading, page)
            page += 1
        if isinstance(raw.get("notes"), str) and raw["notes"].strip():
            first_slide.notes_slide.notes_text_frame.text = raw["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
