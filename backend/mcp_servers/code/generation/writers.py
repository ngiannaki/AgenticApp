"""Renders normalized content (documents.py / tables.py) into real file
bytes for each supported format.

"Real formatting" means actual document/spreadsheet structure — Word
heading styles and table borders via python-docx, bold shaded header rows
and auto-sized columns via openpyxl — not just text dumped into a file
with the right extension.
"""

from __future__ import annotations

import csv
import io
from typing import Any

import openpyxl
from docx import Document
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.mcp_servers.code.generation.documents import normalize_document
from backend.mcp_servers.code.generation.tables import normalize_table_content

# --- txt -------------------------------------------------------------------


def render_txt(content: dict[str, Any]) -> bytes:
    doc = normalize_document(content)
    lines: list[str] = []

    if doc["title"]:
        lines.append(doc["title"])
        lines.append("=" * len(doc["title"]))
        lines.append("")

    for section in doc["sections"]:
        if section["heading"]:
            lines.append(section["heading"])
            lines.append("-" * len(section["heading"]))
            lines.append("")
        for p in section["paragraphs"]:
            lines.append(p)
            lines.append("")
        for b in section["bullets"]:
            lines.append(f"  - {b}")
        if section["bullets"]:
            lines.append("")
        for i, n in enumerate(section["numbered"], start=1):
            lines.append(f"  {i}. {n}")
        if section["numbered"]:
            lines.append("")
        if section["table"]:
            lines.extend(_render_text_table(section["table"]))
            lines.append("")

    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def _render_text_table(table: dict[str, Any]) -> list[str]:
    headers = table["headers"]
    rows = table["rows"]
    all_rows = ([headers] if headers else []) + rows
    widths = [max(len(r[c]) for r in all_rows) for c in range(len(all_rows[0]))]

    def fmt_row(row: list[str]) -> str:
        return "  ".join(cell.ljust(widths[c]) for c, cell in enumerate(row))

    lines = []
    if headers:
        lines.append(fmt_row(headers))
        lines.append("  ".join("-" * w for w in widths))
    for row in rows:
        lines.append(fmt_row(row))
    return lines


# --- md ----------------------------------------------------------------------


def render_md(content: dict[str, Any]) -> bytes:
    doc = normalize_document(content)
    lines: list[str] = []

    # Markdown only has one heading namespace (# depth), unlike docx's separate
    # Title style — so nest section headings one level under the title, or
    # they'd render as siblings of it (two H1s) instead of children.
    depth_offset = 1 if doc["title"] else 0
    if doc["title"]:
        lines.append(f"# {doc['title']}")
        lines.append("")

    for section in doc["sections"]:
        if section["heading"]:
            level = min(section["level"] + depth_offset, 6)
            lines.append(f"{'#' * level} {section['heading']}")
            lines.append("")
        for p in section["paragraphs"]:
            lines.append(p)
            lines.append("")
        for b in section["bullets"]:
            lines.append(f"- {b}")
        if section["bullets"]:
            lines.append("")
        for i, n in enumerate(section["numbered"], start=1):
            lines.append(f"{i}. {n}")
        if section["numbered"]:
            lines.append("")
        if section["table"]:
            lines.extend(_render_md_table(section["table"]))
            lines.append("")

    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def _render_md_table(table: dict[str, Any]) -> list[str]:
    headers = table["headers"] or [f"Column {i + 1}" for i in range(len(table["rows"][0]))]

    def esc(cell: str) -> str:
        return cell.replace("|", "\\|").replace("\n", " ")

    lines = ["| " + " | ".join(esc(h) for h in headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in table["rows"]:
        lines.append("| " + " | ".join(esc(c) for c in row) + " |")
    return lines


# --- docx --------------------------------------------------------------------


def render_docx(content: dict[str, Any]) -> bytes:
    doc_content = normalize_document(content)
    document = Document()

    if doc_content["title"]:
        document.add_heading(doc_content["title"], level=0)

    for section in doc_content["sections"]:
        if section["heading"]:
            document.add_heading(section["heading"], level=section["level"])
        for p in section["paragraphs"]:
            document.add_paragraph(p)
        for b in section["bullets"]:
            document.add_paragraph(b, style="List Bullet")
        for n in section["numbered"]:
            document.add_paragraph(n, style="List Number")
        if section["table"]:
            _add_docx_table(document, section["table"])

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _add_docx_table(document: Document, table: dict[str, Any]) -> None:
    headers = table["headers"]
    rows = table["rows"]
    n_cols = len(headers) if headers else len(rows[0])

    doc_table = document.add_table(rows=0, cols=n_cols)
    doc_table.style = "Table Grid"

    if headers:
        header_cells = doc_table.add_row().cells
        for cell, text in zip(header_cells, headers):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    for row in rows:
        cells = doc_table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text


# --- csv -----------------------------------------------------------------------


def render_csv(content: dict[str, Any]) -> bytes:
    sheets = normalize_table_content(content, allow_multi_sheet=False)
    sheet = sheets[0]

    buf = io.StringIO()
    writer = csv.writer(buf)
    if sheet["headers"]:
        writer.writerow(sheet["headers"])
    for row in sheet["rows"]:
        writer.writerow(["" if c is None else c for c in row])

    # utf-8-sig: BOM so Excel opens non-ASCII text correctly instead of mangling it.
    return buf.getvalue().encode("utf-8-sig")


# --- xlsx --------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def render_xlsx(content: dict[str, Any]) -> bytes:
    sheets = normalize_table_content(content, allow_multi_sheet=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for sheet in sheets:
        ws = workbook.create_sheet(sheet["name"])
        _write_xlsx_sheet(ws, sheet)

    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _write_xlsx_sheet(ws: Any, sheet: dict[str, Any]) -> None:
    headers = sheet["headers"]
    rows = sheet["rows"]
    start_row = 1

    if headers:
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=text)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        ws.freeze_panes = "A2"
        start_row = 2

    for r, row in enumerate(rows, start=start_row):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    n_cols = len(headers) if headers else (len(rows[0]) if rows else 0)
    width_rows = ([headers] if headers else []) + [["" if v is None else v for v in row] for row in rows]
    for col in range(1, n_cols + 1):
        widest = max((len(str(r[col - 1])) for r in width_rows), default=10)
        ws.column_dimensions[get_column_letter(col)].width = min(max(widest + 2, 8), 60)


# --- html ----------------------------------------------------------------------
# A self-contained, styled report page (no scripts, no external assets): the
# same document shape as md/docx, plus optional KPI cards
# (`content.metrics`) and simple bar charts per section (`section.chart`),
# drawn as inline SVG. Light/dark aware and print-friendly.

_HTML_CSS = """
:root{--bg:#f6f7f9;--card:#fff;--fg:#16181d;--muted:#667085;--line:#e4e7ec;--accent:#2563eb;--good:#16a34a;--bad:#dc2626}
@media (prefers-color-scheme:dark){:root{--bg:#111318;--card:#1a1d24;--fg:#eceef2;--muted:#98a2b3;--line:#2a2f3a;--accent:#5b93ff;--good:#22c55e;--bad:#f87171}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif}
main{max-width:1100px;margin:0 auto;padding:40px 24px 64px}
h1{font-size:28px;letter-spacing:-.02em;margin:0 0 6px}.subtitle{color:var(--muted);margin:0 0 28px}
h2{font-size:19px;margin:0 0 12px}h3{font-size:16px;margin:0 0 10px}h4{font-size:14px;margin:0 0 8px}
.kpis{display:grid;grid-template-columns:repeat(auto-fill,minmax(200px,1fr));gap:14px;margin:0 0 24px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:16px 18px}
.kpi .label{font-size:12px;font-weight:600;letter-spacing:.04em;text-transform:uppercase;color:var(--muted)}
.kpi .value{font-size:26px;font-weight:700;letter-spacing:-.02em;margin-top:4px;font-variant-numeric:tabular-nums}
.kpi .note{font-size:13px;color:var(--muted)}.kpi .note.up{color:var(--good)}.kpi .note.down{color:var(--bad)}
section{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:22px 24px;margin:0 0 18px}
p{margin:0 0 10px}ul,ol{margin:0 0 10px;padding-left:22px}li{margin:2px 0}
.table-wrap{overflow-x:auto;margin-top:8px}table{border-collapse:collapse;width:100%;font-size:14px}
th{text-align:left;font-size:12px;letter-spacing:.04em;text-transform:uppercase;color:var(--muted);border-bottom:1px solid var(--line);padding:8px 12px;white-space:nowrap}
td{border-bottom:1px solid var(--line);padding:9px 12px}tr:last-child td{border-bottom:none}td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
.chart{margin:8px 0 6px}.chart text{fill:var(--muted);font-size:12px}.legend{display:flex;gap:14px;flex-wrap:wrap;font-size:13px;color:var(--muted);margin-top:4px}
.legend span::before{content:"";display:inline-block;width:10px;height:10px;border-radius:3px;margin-right:6px;background:var(--c)}
footer{color:var(--muted);font-size:12px;margin-top:24px}
@media print{body{background:#fff}section,.kpi{break-inside:avoid;border-color:#ddd}}
"""
_PALETTE = ["#2563eb", "#16a34a", "#f59e0b", "#9333ea", "#0891b2", "#dc2626"]


def _esc(text: Any) -> str:
    import html as _html

    return _html.escape("" if text is None else str(text))


def _is_num(cell: str) -> bool:
    s = cell.replace(",", "").replace("$", "").replace("%", "").replace("€", "").strip()
    if s.startswith("(") and s.endswith(")"):
        s = s[1:-1]
    try:
        float(s)
        return True
    except ValueError:
        return False


def _html_table(table: dict[str, Any]) -> str:
    rows = table["rows"]
    headers = table["headers"] or []
    width = len(headers) if headers else len(rows[0])
    numeric = [bool(rows) and all(_is_num(r[i]) or r[i] in ("", "—", "-") for r in rows) for i in range(width)]
    head = "".join(f'<th class="{"num" if numeric[i] else ""}">{_esc(h)}</th>' for i, h in enumerate(headers))
    body = "".join(
        "<tr>" + "".join(f'<td class="{"num" if numeric[i] else ""}">{_esc(c)}</td>' for i, c in enumerate(r)) + "</tr>"
        for r in rows
    )
    return f'<div class="table-wrap"><table>{f"<thead><tr>{head}</tr></thead>" if head else ""}<tbody>{body}</tbody></table></div>'


def _html_chart(chart: Any) -> str:
    """{"title"?, "labels": [...], "series": [{"name", "values": [...]}]} -> grouped bar chart."""
    if not isinstance(chart, dict):
        return ""
    labels = [str(x) for x in (chart.get("labels") or [])][:24]
    series = [s for s in (chart.get("series") or []) if isinstance(s, dict)][:6]
    if not labels or not series:
        return ""
    values = [[(float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0) for v in (s.get("values") or [])][: len(labels)] for s in series]
    top = max([max(v, default=0.0) for v in values] + [0.0]) or 1.0
    w, h, pad_l, pad_b, pad_t = 760, 260, 8, 28, 16
    group = (w - pad_l) / len(labels)
    bar = max(4.0, min(36.0, (group * 0.7) / len(series)))
    parts = [f'<line x1="{pad_l}" y1="{h - pad_b}" x2="{w}" y2="{h - pad_b}" stroke="currentColor" opacity=".2"/>']
    for i, label in enumerate(labels):
        x0 = pad_l + i * group + (group - bar * len(series)) / 2
        for j, vals in enumerate(values):
            v = vals[i] if i < len(vals) else 0.0
            bh = max(0.0, (h - pad_b - pad_t) * v / top)
            x, y = x0 + j * bar, h - pad_b - bh
            parts.append(
                f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar - 2:.1f}" height="{bh:.1f}" rx="3" fill="{_PALETTE[j % len(_PALETTE)]}">'
                f'<title>{_esc(series[j].get("name"))} · {_esc(label)}: {v:,.2f}</title></rect>'
            )
        parts.append(f'<text x="{pad_l + i * group + group / 2:.1f}" y="{h - 8}" text-anchor="middle">{_esc(label)}</text>')
    legend = "".join(f'<span style="--c:{_PALETTE[j % len(_PALETTE)]}">{_esc(s.get("name"))}</span>' for j, s in enumerate(series))
    title = f"<h4>{_esc(chart.get('title'))}</h4>" if chart.get("title") else ""
    return (f'<div class="chart">{title}<svg viewBox="0 0 {w} {h}" width="100%" role="img" aria-label="{_esc(chart.get("title") or "chart")}">'
            + "".join(parts) + f'</svg><div class="legend">{legend}</div></div>')


def render_html(content: dict[str, Any]) -> bytes:
    doc = normalize_document(content)
    raw_sections = content.get("sections") or []
    out = ['<!doctype html><html lang="en"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           f"<title>{_esc(doc['title'] or 'Report')}</title><style>{_HTML_CSS}</style></head><body><main>"]
    if doc["title"]:
        out.append(f"<h1>{_esc(doc['title'])}</h1>")
    if isinstance(content.get("subtitle"), str):
        out.append(f'<p class="subtitle">{_esc(content["subtitle"])}</p>')
    metrics = [m for m in (content.get("metrics") or []) if isinstance(m, dict)][:12]
    if metrics:
        out.append('<div class="kpis">')
        for m in metrics:
            note = m.get("note") or m.get("change") or ""
            # Coloured by sign, unless `trend` says whether the change is good ("good"/"bad").
            trend = {"good": "up", "bad": "down"}.get(str(m.get("trend") or "").lower())
            if trend is None:
                trend = "up" if str(note).strip().startswith("+") else "down" if str(note).strip().startswith(("-", "−")) else ""
            out.append(f'<div class="kpi"><div class="label">{_esc(m.get("label"))}</div><div class="value">{_esc(m.get("value"))}</div>'
                       + (f'<div class="note {trend}">{_esc(note)}</div>' if note else "") + "</div>")
        out.append("</div>")
    for i, s in enumerate(doc["sections"]):
        out.append("<section>")
        if s["heading"]:
            tag = f"h{min(s['level'] + 1, 4)}"
            out.append(f"<{tag}>{_esc(s['heading'])}</{tag}>")
        out.extend(f"<p>{_esc(p)}</p>" for p in s["paragraphs"])
        if s["bullets"]:
            out.append("<ul>" + "".join(f"<li>{_esc(b)}</li>" for b in s["bullets"]) + "</ul>")
        if s["numbered"]:
            out.append("<ol>" + "".join(f"<li>{_esc(n)}</li>" for n in s["numbered"]) + "</ol>")
        raw = raw_sections[i] if i < len(raw_sections) and isinstance(raw_sections[i], dict) else {}
        out.append(_html_chart(raw.get("chart")))
        if s["table"]:
            out.append(_html_table(s["table"]))
        out.append("</section>")
    if isinstance(content.get("footer"), str):
        out.append(f"<footer>{_esc(content['footer'])}</footer>")
    out.append("</main></body></html>")
    return "".join(out).encode("utf-8")
