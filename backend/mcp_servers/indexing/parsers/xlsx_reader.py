from pathlib import Path

import openpyxl

# Legacy binary (.xls) and OpenDocument (.ods) spreadsheets aren't zip/xml
# in openpyxl's expected shape, so it raises InvalidFileException on both —
# they need a different engine via pandas instead.
_LEGACY_ENGINES = {".xls": "xlrd", ".ods": "odf"}


def parse_xlsx(xlsx_path: str | Path) -> dict:
    """Parse an Excel workbook into one entry per sheet (headers + rows).

    `read_only=True` streams rows instead of loading the whole workbook into
    memory, and `data_only=True` reads formulas' last-calculated values
    rather than the formula text itself.
    """

    xlsx_path = Path(xlsx_path)

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

    engine = _LEGACY_ENGINES.get(xlsx_path.suffix.lower())
    if engine is not None:
        return _parse_legacy(xlsx_path, engine)

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    try:
        sheets = []
        markdown_parts = []
        text_parts = []

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            all_rows = [
                ["" if cell is None else cell for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]

            headers = all_rows[0] if all_rows else []
            rows = all_rows[1:]

            sheets.append({
                "name": sheet_name,
                "row_count": len(rows),
                "column_count": len(headers),
                "headers": headers,
                "rows": rows,
            })

            markdown_parts.append(f"## {sheet_name}\n" + _to_markdown_table(headers, rows))
            text_parts.append(f"=== {sheet_name} ===\n" + _to_plain_text(headers, rows))
    finally:
        workbook.close()

    return {
        "filename": xlsx_path.name,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "markdown": "\n\n".join(markdown_parts),
        "text": "\n\n".join(text_parts),
    }


def _parse_legacy(xlsx_path: Path, engine: str) -> dict:
    """Parse a legacy `.xls` (engine="xlrd") or OpenDocument `.ods`
    (engine="odf") spreadsheet via pandas, since openpyxl can't open either.
    Both are soft dependencies — only imported here, on this path."""
    import pandas as pd

    all_sheets = pd.read_excel(xlsx_path, sheet_name=None, engine=engine, header=None)

    sheets = []
    markdown_parts = []
    text_parts = []

    for sheet_name, df in all_sheets.items():
        all_rows = df.fillna("").values.tolist()
        headers = all_rows[0] if all_rows else []
        rows = all_rows[1:]

        sheets.append({
            "name": sheet_name,
            "row_count": len(rows),
            "column_count": len(headers),
            "headers": headers,
            "rows": rows,
        })

        markdown_parts.append(f"## {sheet_name}\n" + _to_markdown_table(headers, rows))
        text_parts.append(f"=== {sheet_name} ===\n" + _to_plain_text(headers, rows))

    return {
        "filename": xlsx_path.name,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "markdown": "\n\n".join(markdown_parts),
        "text": "\n\n".join(text_parts),
    }


def _to_markdown_table(headers: list, rows: list[list]) -> str:
    if not headers:
        return "\n".join(" | ".join(str(c) for c in row) for row in rows)

    lines = [
        " | ".join(str(h) for h in headers),
        " | ".join("---" for _ in headers),
    ]
    lines.extend(" | ".join(str(c) for c in row) for row in rows)
    return "\n".join(lines)


def _to_plain_text(headers: list, rows: list[list]) -> str:
    lines = [", ".join(str(h) for h in headers)] if headers else []
    lines.extend(", ".join(str(c) for c in row) for row in rows)
    return "\n".join(lines)
