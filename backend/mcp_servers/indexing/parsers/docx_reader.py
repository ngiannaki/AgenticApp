from pathlib import Path

import docx


def parse_docx(docx_path: str | Path) -> dict:
    """Parse a Word document into full text + heading structure + a markdown rendering.

    Paragraphs are walked in document order; a paragraph whose style is
    "Heading N" becomes a markdown `#`*N heading and is also recorded in
    `headings` (so callers can chunk by section without re-parsing). Tables
    are rendered inline, at their position in the document, as markdown
    tables — same convention as the CSV/XLSX readers.
    """

    docx_path = Path(docx_path)

    if not docx_path.exists():
        raise FileNotFoundError(f"DOCX not found: {docx_path}")

    document = docx.Document(str(docx_path))

    headings: list[dict] = []
    markdown_lines: list[str] = []
    text_lines: list[str] = []
    char_offset = 0

    # python-docx exposes paragraphs and tables as separate flat lists with no
    # shared ordering, so walk the underlying document body directly to
    # interleave them in the order they actually appear.
    body = document.element.body
    table_iter = iter(document.tables)
    paragraph_iter = iter(document.paragraphs)

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = next(paragraph_iter, None)
            if paragraph is None:
                continue
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = (paragraph.style.name or "") if paragraph.style else ""
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
                level = max(1, min(level, 6))
                headings.append({"level": level, "text": text, "char_offset": char_offset})
                markdown_lines.append(f"{'#' * level} {text}")
            else:
                markdown_lines.append(text)

            text_lines.append(text)
            char_offset += len(text) + 1

        elif tag == "tbl":
            table = next(table_iter, None)
            if table is None:
                continue
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers, data_rows = rows[0], rows[1:]
            markdown_lines.append(_to_markdown_table(headers, data_rows))
            table_text = _to_plain_text(headers, data_rows)
            text_lines.append(table_text)
            char_offset += len(table_text) + 1

    full_text = "\n".join(text_lines)

    return {
        "filename": docx_path.name,
        "full_text": full_text,
        "headings": headings,
        "heading_count": len(headings),
        "markdown": "\n\n".join(markdown_lines),
        "text": full_text,
    }


def _to_markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    lines = [
        " | ".join(headers),
        " | ".join("---" for _ in headers),
    ]
    lines.extend(" | ".join(row) for row in rows)
    return "\n".join(lines)


def _to_plain_text(headers: list[str], rows: list[list[str]]) -> str:
    lines = [", ".join(headers)]
    lines.extend(", ".join(row) for row in rows)
    return "\n".join(lines)
