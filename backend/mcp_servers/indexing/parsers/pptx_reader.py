from pathlib import Path

from pptx import Presentation


def parse_pptx(pptx_path: str | Path) -> dict:
    """Parse a PowerPoint deck into one "page" per slide.

    Shaped like parse_pdf's result ({filename, page_count, pages, full_text})
    so it can reuse the same page-based/hierarchical chunking path as PDFs —
    a slide here plays the same role a page does there. Each slide's text
    combines every text-frame shape (title first, if present) plus its
    speaker notes, since notes often carry context a slide's bullet points
    alone don't.
    """

    pptx_path = Path(pptx_path)

    if not pptx_path.exists():
        raise FileNotFoundError(f"PPTX not found: {pptx_path}")

    presentation = Presentation(str(pptx_path))

    pages = []
    full_text_parts = []

    for i, slide in enumerate(presentation.slides, start=1):
        title = ""
        body_parts = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape == slide.shapes.title:
                title = text
            else:
                body_parts.append(text)

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        parts = ([title] if title else []) + body_parts
        if notes:
            parts.append(f"[Notes: {notes}]")

        slide_text = "\n".join(parts)

        pages.append({"page_number": i, "slide_number": i, "title": title, "text": slide_text})
        full_text_parts.append(f"## Slide {i}" + (f": {title}" if title else "") + f"\n{slide_text}")

    return {
        "filename": pptx_path.name,
        "page_count": len(pages),
        "pages": pages,
        "full_text": "\n\n".join(full_text_parts),
    }
